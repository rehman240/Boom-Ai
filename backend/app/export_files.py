"""Draw a campaign package (see app/review.py) as a PDF or an editable Word file.

Both files show the same blocks in the same order, so they always say the same thing.
The PDF uses the bundled DejaVu Sans font (app/fonts, free licence in LICENSE_DEJAVU), which
covers Latin, Greek and Cyrillic text. The Word file uses the reader's own fonts.
"""

import io
import re
import unicodedata
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt, RGBColor
from fpdf import FPDF
from fpdf.fonts import FontFace

from app.brand import APP_NAME
from app.review import APPROVED, MISSING

FONTS = Path(__file__).parent / "fonts"
STATUS_TEXT = {APPROVED: "Approved", "waiting": "Not approved yet", MISSING: "Not done yet"}

INK = (24, 24, 32)
MUTED = (100, 104, 120)
ACCENT = (91, 63, 214)


def _meta(pkg: dict[str, Any]) -> str:
    when = pkg["generated_at"].strftime("%b %d, %Y, %H:%M UTC").replace(" 0", " ")
    return f"{pkg['version_label']}  ·  Exported {when}"


def filename(pkg: dict[str, Any], ext: str) -> str:
    """An ASCII file name, e.g. "NOVA-Desk-Lamp-version-12.pdf"."""
    name = unicodedata.normalize("NFKD", pkg["title"]).encode("ascii", "ignore").decode()
    name = re.sub(r"[^A-Za-z0-9]+", "-", name).strip("-")[:60] or "campaign"
    version = re.sub(r"\D", "", pkg["version_label"])
    return f"{name}-version-{version}.{ext}"


# --- PDF ----------------------------------------------------------------------------------


class _Pdf(FPDF):
    def __init__(self, pkg: dict[str, Any]):
        super().__init__(format="A4")
        self.pkg = pkg
        self.add_font("DejaVu", "", FONTS / "DejaVuSans.ttf")
        self.add_font("DejaVu", "B", FONTS / "DejaVuSans-Bold.ttf")
        self.set_margins(18, 18, 18)
        self.set_auto_page_break(True, margin=18)
        self.set_title(pkg["title"])
        self.set_author(APP_NAME)

    def footer(self) -> None:
        self.set_y(-12)
        self.set_font("DejaVu", "", 7.5)
        self.set_text_color(*MUTED)
        self.cell(0, 5, f"{self.pkg['title']}  ·  {self.pkg['version_label']}  ·  Not published", align="L")
        self.set_x(self.l_margin)
        self.cell(0, 5, f"Page {self.page_no()}", align="R")

    def text_block(self, text: str, size: float = 9.5, bold: bool = False, color=INK, after: float = 1.5) -> None:
        self.set_font("DejaVu", "B" if bold else "", size)
        self.set_text_color(*color)
        self.multi_cell(0, size * 0.5, text, align="L", new_x="LMARGIN", new_y="NEXT")
        self.ln(after)

    def notice(self, text: str, color: tuple[int, int, int]) -> None:
        self.set_draw_color(*color)
        self.set_text_color(*color)
        self.set_font("DejaVu", "B", 9)
        self.multi_cell(0, 5, text, border=1, padding=(2, 3), align="L", new_x="LMARGIN", new_y="NEXT")
        self.ln(2)

    def grid(self, headers: list[str], rows: list[list[str]], widths: tuple[int, ...] | None = None) -> None:
        self.set_font("DejaVu", "", 8.5)
        self.set_text_color(*INK)
        self.set_draw_color(200, 200, 210)
        heading = FontFace(emphasis="BOLD", fill_color=(238, 236, 250))
        with self.table(col_widths=widths, headings_style=heading, line_height=5, padding=1.5) as table:
            table.row(headers)
            for row in rows:
                table.row([str(c) for c in row])
        self.ln(2)


def to_pdf(pkg: dict[str, Any]) -> bytes:
    pdf = _Pdf(pkg)
    pdf.add_page()
    pdf.text_block(f"{APP_NAME}  ·  Campaign package", 9, color=ACCENT, after=1)
    pdf.text_block(pkg["title"], 20, bold=True, after=1)
    pdf.text_block(_meta(pkg), 8.5, color=MUTED, after=4)
    pdf.notice(pkg["not_published"], (180, 60, 40))
    pdf.notice(pkg["status_note"], (30, 130, 80) if pkg["final"] else (150, 110, 10))
    pdf.text_block(pkg["ai_note"], 8.5, color=MUTED, after=4)

    pdf.text_block("Approval checklist", 13, bold=True, after=2)
    pdf.grid(["Part", "Status", "Detail"],
             [[r["label"], STATUS_TEXT[r["status"]], r["detail"]] for r in pkg["checklist"]], (30, 25, 45))

    for section in pkg["sections"]:
        pdf.ln(3)
        pdf.text_block(section["heading"], 13, bold=True, color=ACCENT, after=2)
        for block in section["blocks"]:
            kind = block[0]
            if kind == "p":
                pdf.text_block(block[1], 9, color=MUTED)
            elif kind == "sub":
                pdf.ln(1)
                pdf.text_block(block[1], 10.5, bold=True, after=1)
            elif kind == "field":
                pdf.text_block(block[1], 8, bold=True, color=MUTED, after=0.5)
                pdf.text_block(block[2], 9.5, after=2)
            elif kind == "list":
                if block[1]:
                    pdf.text_block(block[1], 8, bold=True, color=MUTED, after=0.5)
                for entry in block[2]:
                    pdf.text_block(f"•  {entry}", 9.5, after=0.8)
                pdf.ln(1.2)
            elif kind == "table":
                pdf.grid(block[1], block[2])
    return bytes(pdf.output())


# --- Word ---------------------------------------------------------------------------------


def _small(paragraph, text: str, *, bold: bool = False, color: tuple[int, int, int] = MUTED, size: float = 9) -> None:
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(*color)


def _table(doc, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True
    for row in rows:
        for cell, text in zip(table.add_row().cells, row):
            cell.text = str(text)
    doc.add_paragraph()


def to_docx(pkg: dict[str, Any]) -> bytes:
    doc = Document()
    doc.core_properties.title = pkg["title"]
    doc.core_properties.author = APP_NAME
    doc.styles["Normal"].font.size = Pt(10.5)

    _small(doc.add_paragraph(), f"{APP_NAME}  ·  Campaign package", color=ACCENT)
    doc.add_heading(pkg["title"], level=0)
    _small(doc.add_paragraph(), _meta(pkg))
    _small(doc.add_paragraph(), pkg["not_published"], bold=True, color=(180, 60, 40), size=10.5)
    _small(doc.add_paragraph(), pkg["status_note"], bold=True,
           color=(30, 130, 80) if pkg["final"] else (150, 110, 10), size=10.5)
    _small(doc.add_paragraph(), pkg["ai_note"])

    doc.add_heading("Approval checklist", level=1)
    _table(doc, ["Part", "Status", "Detail"],
           [[r["label"], STATUS_TEXT[r["status"]], r["detail"]] for r in pkg["checklist"]])

    for section in pkg["sections"]:
        doc.add_heading(section["heading"], level=1)
        for block in section["blocks"]:
            kind = block[0]
            if kind == "p":
                _small(doc.add_paragraph(), block[1])
            elif kind == "sub":
                doc.add_heading(block[1], level=2)
            elif kind == "field":
                p = doc.add_paragraph()
                _small(p, block[1], bold=True, size=9)
                p.paragraph_format.space_after = Pt(0)
                doc.add_paragraph(block[2])
            elif kind == "list":
                if block[1]:
                    p = doc.add_paragraph()
                    _small(p, block[1], bold=True, size=9)
                    p.paragraph_format.space_after = Pt(0)
                for entry in block[2]:
                    doc.add_paragraph(entry, style="List Bullet")
            elif kind == "table":
                _table(doc, block[1], block[2])

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
