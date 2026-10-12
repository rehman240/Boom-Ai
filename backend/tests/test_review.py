"""Review and Export: the approval checklist and the PDF and Word files (client brief 4.9)."""

import io

from docx import Document
from sqlalchemy import select

from app import export_files, review
from app.models import Event, Project, ProjectStatus

from tests.test_assets import generate as generate_assets
from tests.test_assets import stage as asset_stage
from tests.test_assets import with_direction
from tests.test_audiences import confirmed_project
from tests.test_brief_summary import signup
from tests.test_budget import generate as generate_budget
from tests.test_budget import stage as budget_stage


def checklist(client, project_id):
    r = client.get(f"/projects/{project_id}/review")
    assert r.status_code == 200, r.text
    return r.json()


def rows(state):
    return {r["key"]: r for r in state["checklist"]}


def complete(client):
    """A campaign with every part made: brief, audience, direction, assets, budget, measurement."""
    project_id = with_direction(client)
    generate_assets(client, project_id)
    generate_budget(client, project_id)
    assert client.post(f"/projects/{project_id}/conversions/plan").status_code == 201
    return project_id


def approve_all(client, project_id):
    for row in checklist(client, project_id)["checklist"]:
        for item_id in row["item_ids"]:
            assert client.post(f"/projects/{project_id}/items/{item_id}/approve").status_code == 200


def docx_text(body: bytes) -> str:
    doc = Document(io.BytesIO(body))
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


def test_the_checklist_names_the_six_parts_and_what_to_do_next(client):
    state = checklist(client, confirmed_project(client))
    assert [r["key"] for r in state["checklist"]] == ["brief", "target", "concept", "assets", "budget", "measurement"]
    by_key = rows(state)
    assert by_key["brief"]["status"] == "approved"
    assert by_key["target"] == {**by_key["target"], "status": "missing", "detail": "Choose a primary audience."}
    assert by_key["assets"]["status"] == "missing"
    assert state["final"] is False
    assert state["not_published"].startswith("Not published.")


def test_approving_every_part_makes_the_campaign_final_and_an_edit_takes_it_back(client, db):
    project_id = complete(client)
    by_key = rows(checklist(client, project_id))
    assert {k: r["status"] for k, r in by_key.items() if k != "brief"} == {
        "target": "waiting", "concept": "waiting", "assets": "waiting", "budget": "waiting", "measurement": "waiting",
    }
    assert by_key["assets"]["detail"] == "0 of 7 assets approved."
    assert len(by_key["assets"]["item_ids"]) == 7

    approve_all(client, project_id)
    state = checklist(client, project_id)
    assert state["final"] is True
    assert state["status_note"].startswith("Final.")
    assert client.get("/projects").json()["stats"]["ready_to_export"] == 1

    # Editing approved text withdraws its approval, so the campaign is no longer ready.
    short_ad = next(a for a in asset_stage(client, project_id)["assets"] if a["key"] == "short_ad")
    client.patch(f"/projects/{project_id}/items/{short_ad['id']}", json={"data": {"headline": "A new headline"}})
    state = checklist(client, project_id)
    assert state["final"] is False
    assert rows(state)["assets"]["detail"] == "6 of 7 assets approved."
    db.expire_all()
    assert db.get(Project, project_id).status == ProjectStatus.IN_PROGRESS
    assert client.get("/projects").json()["stats"]["ready_to_export"] == 0


def test_a_brief_changed_after_confirming_needs_checking_again(client):
    project_id = confirmed_project(client)
    client.patch(f"/projects/{project_id}/brief", json={"differentiators": "Now with a USB-C port."})
    assert rows(checklist(client, project_id))["brief"]["status"] == "waiting"


def test_the_package_holds_every_part_as_it_is_now(client, db):
    project_id = complete(client)
    plan = budget_stage(client, project_id)["plan"]
    pkg = review.package(db, db.get(Project, project_id))
    assert [s["heading"] for s in pkg["sections"]][:6] == [
        "Campaign brief", "Target audience", "Campaign direction", "Creative assets", "Budget", "Measurement plan",
    ]
    assert pkg["version_label"] == f"Version {db.get(Project, project_id).version}"
    assert pkg["final"] is False and pkg["status_note"].startswith("Draft.")

    budget = next(s for s in pkg["sections"] if s["heading"] == "Budget")
    table = next(b for b in budget["blocks"] if b[0] == "table")
    assert [row[0] for row in table[2]] == [line["channel"] for line in plan["data"]["lines"]]
    assert ("field", "Media budget", "$12,000") in budget["blocks"]


def test_word_export_is_editable_and_says_it_is_not_published(client):
    project_id = complete(client)
    short_ad = next(a for a in asset_stage(client, project_id)["assets"] if a["key"] == "short_ad")
    r = client.get(f"/projects/{project_id}/export.docx")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert r.headers["content-disposition"] == 'attachment; filename="NOVA-Desk-Lamp-version-' + \
        r.headers["content-disposition"].split("version-")[1]
    assert r.headers["cache-control"] == "no-store"
    text = docx_text(r.content)
    assert "Not published." in text
    assert "Draft." in text
    assert short_ad["data"]["headline"] in text
    assert "Approval checklist" in text and "Not approved yet" in text


def test_pdf_export_is_a_pdf_and_counts_only_that_it_happened(client, db):
    project_id = complete(client)
    approve_all(client, project_id)
    r = client.get(f"/projects/{project_id}/export.pdf")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF")
    assert r.headers["content-disposition"].endswith('.pdf"')
    names = db.scalars(select(Event.name).where(Event.project_id == project_id)).all()
    assert "exported_pdf" in names


def test_an_early_export_works_and_says_what_is_missing(client):
    project_id = confirmed_project(client)
    for fmt in ("pdf", "docx"):
        assert client.get(f"/projects/{project_id}/export.{fmt}").status_code == 200
    text = docx_text(client.get(f"/projects/{project_id}/export.docx").content)
    assert "No primary audience chosen yet." in text
    assert "Not done yet" in text


def test_text_in_any_script_does_not_break_the_pdf(client, db):
    project_id = confirmed_project(client)
    client.patch(f"/projects/{project_id}", json={"name": "Café “Lumière” — Привет 漢字 🚀"})
    r = client.get(f"/projects/{project_id}/export.pdf")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert export_files.filename({"title": "Café “Lumière” — 漢字", "version_label": "Version 3"}, "pdf") == (
        "Cafe-Lumiere-version-3.pdf"
    )


def test_exports_are_private(client):
    project_id = confirmed_project(client)
    client.post("/auth/logout")
    signup(client, "other@example.com")
    for path in ("review", "export.pdf", "export.docx"):
        assert client.get(f"/projects/{project_id}/{path}").status_code == 404
